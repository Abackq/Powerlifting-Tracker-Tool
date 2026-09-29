from openpyxl import load_workbook
import re
import time
from datetime import datetime, date
from datetime import timedelta

WEEK_PATTERN = re.compile(r"^(?:Week\s+(\d+)|Meet\s+Week)$", re.IGNORECASE)
DAY_PATTERN = re.compile(r"Day\s+(\d+)$", re.IGNORECASE)
DATE_FORMATS = ["%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%B %d, %Y"]

# def print_sheet_grid(workbook, sheet_name: str):
#     """Prints every non-empty cell's row, column, and value for a given sheet.
#     Useful for visually inspecting layout before writing detection logic."""
#     sheet = workbook[sheet_name]
#     for row in sheet.iter_rows():
#         for cell in row:
#             if cell.value is not None:
#                 print(f"row={cell.row}, col={cell.column}, value={cell.value!r}")

# region Stage 1: Loading
def load_client_workbook(filepath: str):
    workbook = load_workbook(filepath, read_only=False, data_only=True)
    return workbook

def split_tabs(workbook, metadata_tab_name: str = "READ ME"):
    """Partitions a workbook's tabs into (metadata_sheet, block_tab_names).
    metadata_sheet is None if no matching tab exists; the caller decides whether that's an error."""
    metadata_sheet = None
    block_tabs = []
    for name in workbook.sheetnames:
        if name.strip() == metadata_tab_name.strip():
            metadata_sheet = workbook[name]
        else:
            block_tabs.append(name)
    return metadata_sheet, block_tabs

def get_block_tabs(workbook, metadata_tab_name: str = "READ ME") -> list[str]:
    "Returns tab names in order, excluding metadata tabs, preserving order for block numbering"
    return split_tabs(workbook, metadata_tab_name)[1]

# endregion Stage 1: Loading

# region Stage 2: Structure Mapping (weeks, days, blocks)
def find_week_columns(sheet) -> list[dict]:
    "Scans the sheet for 'Week N' labels and returns their row column, and week number."
    weeks = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                match = WEEK_PATTERN.match(cell.value.strip())
                if match:
                    is_meet_week = match.group(1) is None
                    weeks.append({
                        "week_number": len(weeks) + 1 if is_meet_week else int(match.group(1)),
                        "is_meet_week": is_meet_week,
                        "row": cell.row,
                        "start_column": cell.column
                    })
    return weeks

def find_day_rows(sheet, week: dict) -> list[dict]:
    "Scans down from the start column of the given week to find 'Day N' labels."
    days = []
    col = week["start_column"]
    for row in sheet.iter_rows(min_row=week["row"] + 1, min_col=col, max_col=col):
        cell = row[0]
        if isinstance(cell.value, str):
            match = DAY_PATTERN.match(cell.value.strip())
            if match:
                days.append({
                    "day_number": int(match.group(1)),
                    "row": cell.row
                })
    return days

def build_week_day_structure(sheet) -> list[dict]:
    "Combines week detection and day detection into a structured list."
    weeks = find_week_columns(sheet)
    for week in weeks:
        week["days"] = find_day_rows(sheet, week)
        compute_day_ranges(sheet, week)
    return weeks

def build_all_blocks(workbook) -> list[dict]:
    """Builds the full week/day structure for every training block tab, attaching 
    the sequential block number based on tab order."""
    block_tabs = get_block_tabs(workbook)
    blocks = []
    for block_number, tab_name in enumerate(block_tabs, start=1):
        sheet = workbook[tab_name]
        blocks.append({
            "block_number": block_number,
            "tab_name": tab_name,
            "weeks": build_week_day_structure(sheet)
        })
    return blocks

def resolve_day_end(sheet, day: dict, search_limit: int, exercise_col: int) -> None:
    """Scans downward from day['start_row'] until a blank row is found,
    but never searches past search_limit (the next day's row, or sheet's max_row)."""
    row = day["start_row"]
    while row < search_limit:
        cell = sheet.cell(row=row, column=exercise_col)
        if cell.value is None:
            day["end_row"] = row
            return
        row += 1
    day["end_row"] = search_limit  # reached the limit without finding a blank row

def compute_day_ranges(sheet, week: dict) -> list[dict]:
    """Adds start_row and end_row to each day, using blank-row detection bounded
    by the next day's row (or sheet end, for the last day)."""
    days = week["days"]
    exercise_col = week["start_column"]
    max_row = sheet.max_row
    for i, day in enumerate(days):
        day["start_row"] = day["row"] + 1  # Exercises start the row after "Day N"
        search_limit = days[i + 1]["row"] if i + 1 < len(days) else max_row + 1
        resolve_day_end(sheet, day, search_limit, exercise_col)
    return days

# endregion Stage 2: Structure Mapping (weeks, days, blocks)

# region Stage 3: Extraction
def parse_day_headers(sheet, week: dict, day: dict) -> dict:
    """Reads the header row for a given day and returns a dict mapping
    field name -> column number, handling the duplicate 'load' header case."""
    header_row = day["row"]
    start_col = week["start_column"]
    max_col = start_col + 10  # generous upper bound for one day's column band; adjust if needed

    load_columns = []
    field_columns = {}

    for col in range(start_col, max_col):
        cell = sheet.cell(row=header_row, column=col)
        if not isinstance(cell.value, str):
            continue
        header = cell.value.strip().lower()
        if header == "load":
            load_columns.append(col)
        elif header in ("rpe", "notes", "range", "top end"):
            field_columns[header] = col

    if len(load_columns) >= 2:
        field_columns["prescribed_rpe_or_backoff"] = load_columns[0]
        field_columns["actual_load"] = load_columns[1]
    elif len(load_columns) == 1:
        field_columns["actual_load"] = load_columns[0]  # only one load column present — no prescribed value to disambiguate

    return field_columns

def extract_day_exercises(sheet, week: dict, day: dict) -> list[dict]:
    """Extracts each exercise row within a day's row range into a raw dict."""
    fields = parse_day_headers(sheet, week, day)
    name_col = week["start_column"]
    sets_reps_col = week["start_column"] + 1

    exercises = []
    for row in range(day["start_row"], day["end_row"]):
        name_cell = sheet.cell(row=row, column=name_col)
        if name_cell.value is None:
            continue  # skip any stray blank row within the range

        exercise = {
            "block_number": None,  # filled in by caller, which knows the block
            "week_number": week["week_number"],
            "day_number": day["day_number"],
            "row": row,
            "exercise_name": name_cell.value,
            "sets_x_reps": sheet.cell(row=row, column=sets_reps_col).value,
            "prescribed_raw": sheet.cell(row=row, column=fields["prescribed_rpe_or_backoff"]).value if "prescribed_rpe_or_backoff" in fields else None,
            "actual_load_raw": sheet.cell(row=row, column=fields["actual_load"]).value if "actual_load" in fields else None,
            "actual_rpe_raw": sheet.cell(row=row, column=fields["rpe"]).value if "rpe" in fields else None,
            "notes": sheet.cell(row=row, column=fields["notes"]).value if "notes" in fields else None,
        }
        exercises.append(exercise)
    return exercises

def extract_all_exercises(workbook, all_blocks: list[dict], client_id) -> list[dict]:
    """Walks every block -> week -> day, calls extract_day_exercises for each,
    and flattens the results into one single list of raw exercise rows."""
    all_raw_rows = []
    for block in all_blocks:
        sheet = workbook[block["tab_name"]]
        for week in block["weeks"]:
            for day in week["days"]:
                rows = extract_day_exercises(sheet, week, day)
                for row in rows:
                    row["block_number"] = block["block_number"]
                    row["client_id"] = client_id
                all_raw_rows.extend(rows)
    return all_raw_rows

# endregion Stage 3: Extraction

# region Stage 4: Validation
def split_prescribed_value(prescribed_raw) -> dict:
    """Splits the ambiguous prescribed column into either a top-set RPE
    or a backdown percentage, based on sign.
    Raises ValueError if prescribed_raw is present but not numeric."""
    if prescribed_raw is None:
        return {"prescribed_rpe": None, "prescribed_backoff_pct": None, "set_type": None}
    if not isinstance(prescribed_raw, (int, float)):
        raise ValueError(f"prescribed value is not numeric: {prescribed_raw!r}")
    if prescribed_raw >= 0:
        return {"prescribed_rpe": prescribed_raw, "prescribed_backoff_pct": None, "set_type": "top"}
    else:
        return {"prescribed_rpe": None, "prescribed_backoff_pct": abs(prescribed_raw) * 100, "set_type": "backdown"}

def validate_numeric_field(value, field_name: str) -> float:
    """Confirms a value (actual_load_raw or actual_rpe_raw) is numeric.
    Raises ValueError if it isn't."""
    if not isinstance(value, (int, float)):
        raise ValueError(f"{field_name} is not numeric: {value!r}")
    return value

def process_exercise_row(raw_row: dict) -> dict:
    """Takes a raw extracted row and returns a clean, validated record.
    Raises ValueError if any field fails validation — the caller is
    responsible for catching this and routing to ingestion_rejects."""
    prescribed = split_prescribed_value(raw_row["prescribed_raw"])
    actual_load = validate_numeric_field(raw_row["actual_load_raw"], "actual_load")
    actual_rpe = validate_numeric_field(raw_row["actual_rpe_raw"], "actual_rpe")

    return {
        "client_id": raw_row["client_id"],
        "block_number": raw_row["block_number"],
        "week_number": raw_row["week_number"],
        "day_number": raw_row["day_number"],
        "exercise_name": raw_row["exercise_name"],
        "sets_x_reps": raw_row["sets_x_reps"],
        "set_type": prescribed["set_type"],
        "prescribed_rpe": prescribed["prescribed_rpe"],
        "prescribed_backoff_pct": prescribed["prescribed_backoff_pct"],
        "actual_load": actual_load,
        "actual_rpe": actual_rpe,
        "notes": raw_row["notes"],
    }

def process_all_exercises(raw_rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Processes every raw extracted row, separating successes from failures.
    Returns (clean_records, rejected_records)."""
    clean_records = []
    rejected_records = []

    for raw_row in raw_rows:
        try:
            clean = process_exercise_row(raw_row)
            clean_records.append(clean)
        except ValueError as e:
            rejected_records.append({
                "client_id": raw_row["client_id"],
                "block_number": raw_row["block_number"],
                "week_number": raw_row["week_number"],
                "day_number": raw_row["day_number"],
                "rejection_reason": str(e)
            })

    return clean_records, rejected_records

# endregion Stage 4: Validation

# region Stage 5: Date Derivation
def find_date_started(workbook) -> date:
    """Finds the 'Date Started' label on the metadata tab and returns the date
    in the cell immediately to its right. Raises ValueError if it can't be found or parsed."""
    metadata_sheet, _ = split_tabs(workbook)
    if metadata_sheet is None:
        raise ValueError("metadata tab not found")

    for row in metadata_sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().lower().startswith("date started"):
                value = metadata_sheet.cell(row=cell.row, column=cell.column + 1).value
                return parse_date_value(value)

    raise ValueError("'Date Started' label not found on metadata tab")

def parse_date_value(value) -> date:
    """Converts a cell value into a date, handling real Excel dates and date-like text."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    raise ValueError(f"could not parse date from: {value!r}")

def compute_block_start_dates(all_blocks: list[dict], date_started: date) -> None:
    """Attaches 'start_date' to each block. Blocks run back-to-back in whole
    Mon-Sun weeks, so each block starts (weeks in previous block) * 7 days after the last."""
    if date_started.weekday() != 0:  # weekday() returns 0 for Monday
        raise ValueError(f"Date Started {date_started} is not a Monday")

    block_start = date_started
    for block in all_blocks:
        block["start_date"] = block_start
        block_length_weeks = max((w["week_number"] for w in block["weeks"]), default=0)
        block_start += timedelta(weeks=block_length_weeks)

def find_short_weeks(all_blocks: list[dict], expected_days: int = 4) -> None:
    """Prints every week whose day count differs from expected_days."""
    for block in all_blocks:
        for week in block["weeks"]:
            day_count = len(week["days"])
            if day_count != expected_days:
                print(block["block_number"], block["tab_name"], "week", week["week_number"], "-", day_count, "days")

def compute_week_date_ranges(all_blocks: list[dict]) -> None:
    """Attaches week_start date and week_end_date to each week, using the block's
    start date and the week's position within its block."""
    for block in all_blocks:
        for week in block["weeks"]:
            week_start = block["start_date"] + timedelta(weeks=week["week_number"] - 1)
            week["week_start_date"] = week_start
            week["week_end_date"] = week_start + timedelta(days=6)

# endregion Stage 5: Date Derivation

if __name__ == "__main__":
    client_id = "paste-jons-actual-client-uuid-here"  # from the clients table in Supabase
    wb = load_client_workbook("Jon Program.xlsx")
    all_blocks = build_all_blocks(wb)
    compute_block_start_dates(all_blocks, find_date_started(wb))
    compute_week_date_ranges(all_blocks)
    raw_rows = extract_all_exercises(wb, all_blocks, client_id)
    clean, rejected = process_all_exercises(raw_rows)
    print(f"{len(clean)} clean, {len(rejected)} rejected")
    # wb = load_client_workbook("Jon Program.xlsx")
    # all_blocks = build_all_blocks(wb)
    # compute_block_start_dates(all_blocks, find_date_started(wb))
    # compute_week_date_ranges(all_blocks)
    # raw_rows = extract_all_exercises(wb, all_blocks)
    # clean, rejected = process_all_exercises(raw_rows)
    # print(f"{len(clean)} clean, {len(rejected)} rejected")
    # print(all_blocks[0]["weeks"][0])
    # wb = load_client_workbook("Jon Program.xlsx")
    # all_blocks = build_all_blocks(wb)
    # print(all_blocks[8])
    # wb = load_client_workbook("Jon Program.xlsx")
    # all_blocks = build_all_blocks(wb)
    # for block in all_blocks:
    #     day_counts = [len(w["days"]) for w in block["weeks"]]
    #     print(block["block_number"], block["tab_name"], day_counts)

    # find_short_weeks(all_blocks)
    # compute_block_start_dates(all_blocks, find_date_started(wb))
    # compute_day_dates(all_blocks)
    # for block in all_blocks[:3]:
    #     print(block["block_number"], block["tab_name"], block["start_date"])
    # print(all_blocks[0]["weeks"][0]["days"])

    # wb = load_client_workbook("Jon Program.xlsx")
    # print(find_date_started(wb))

    # t0 = time.time()
    # wb = load_client_workbook("Jon Program.xlsx")
    # t1 = time.time()
    # all_blocks = build_all_blocks(wb)
    # t2 = time.time()
    # raw_rows = extract_all_exercises(wb, all_blocks)
    # t3 = time.time()
    # clean, rejected = process_all_exercises(raw_rows)
    # t4 = time.time()

    # print(f"{len(clean)} clean, {len(rejected)} rejected")
    # print(f"Load:      {t1 - t0:.2f}s")
    # print(f"Structure: {t2 - t1:.2f}s")
    # print(f"Extract:   {t3 - t2:.2f}s")
    # print(f"Validate:  {t4 - t3:.2f}s")

    # start = time.time()

    # wb = load_client_workbook("Jon Program.xlsx")
    # all_blocks = build_all_blocks(wb)
    # raw_rows = extract_all_exercises(wb, all_blocks)
    # clean, rejected = process_all_exercises(raw_rows)

    # end = time.time()
    # print(f"Clean records: {len(clean)}, Rejected records: {len(rejected)}")
    # print(f"Processing time: {end - start:.2f} seconds")


    # block_tabs = get_block_tabs(wb)
    # sheet = wb[block_tabs[0]]
    # week1 = find_week_columns(sheet)[0]
    # week1["days"] = find_day_rows(sheet, week1)
    # compute_day_ranges(sheet, week1)
    # day1_exercises = extract_day_exercises(sheet, week1, week1["days"][0])
    # for ex in day1_exercises:
    #     print(ex)
        
    # block_tabs = get_block_tabs(wb)
    # week1 = find_week_columns(wb[block_tabs[0]])[0]
    # week1["days"] = find_day_rows(wb[block_tabs[0]], week1)
    # compute_day_ranges(wb[block_tabs[0]], week1)
    # print(week1["days"])

    # all_blocks = build_all_blocks(wb)
    # print(f"Found {len(all_blocks)} blocks")
    # print(all_blocks[0])

    # block_tabs = get_block_tabs(wb)
    # weeks = find_week_columns(wb[block_tabs[0]])
    # days = find_day_rows(wb[block_tabs[0]], weeks[0])
    # weeks[0]["days"] = days
    # compute_day_ranges(weeks[0])
    # print(weeks[0]["days"])

    # INDIVIDUAL FUNCTION TESTS
    # -----------------------------------------------------------------------------
    # block_tabs = get_block_tabs(wb)
    # print(block_tabs)
    # print_sheet_grid(wb, block_tabs[0]) # inspect the first block tab
    # weeks = find_week_columns(wb[block_tabs[0]])
    # days = find_day_rows(wb[block_tabs[0]], weeks[0])
    # structure = build_week_day_structure(wb[block_tabs[0]])
    