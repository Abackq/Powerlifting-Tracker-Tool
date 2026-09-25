from openpyxl import load_workbook
import re

WEEK_PATTERN = re.compile(r"Week\s+(\d+)$", re.IGNORECASE)
DAY_PATTERN = re.compile(r"Day\s+(\d+)$", re.IGNORECASE)

def load_client_workbook(filepath: str):
    workbook = load_workbook(filepath, read_only=True, data_only=True)
    return workbook

def get_block_tabs(workbook, metadata_tab_name: str = "READ ME") -> list[str]:
    "Returns tab names in order, excluding metadata tabs, preserving order for block numbering"
    return [name for name in workbook.sheetnames if name != metadata_tab_name]

def print_sheet_grid(workbook, sheet_name: str):
    """Prints every non-empty cell's row, column, and value for a given sheet.
    Useful for visually inspecting layout before writing detection logic."""
    sheet = workbook[sheet_name]
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                print(f"row={cell.row}, col={cell.column}, value={cell.value!r}")

def find_week_columns(sheet) -> list[dict]:
    "Scans the sheet for 'Week N' labels and returns their row column, and week number."
    weeks = []
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                match = WEEK_PATTERN.match(cell.value.strip())
                if match:
                    weeks.append({
                        "week_number": int(match.group(1)),
                        "row": cell.row,
                        "start_column": cell.column
                    })
    return weeks

def find_day_rows(sheet, week: dict) -> list[dict]:
    "Scans down from the start column of the given week to find 'Day N' labels."
    days = []
    col = week["start_column"]
    for row in sheet.iter_rows(min_row=week["row"] + 1, min_col=col, max_col=col):
        cell = row[0]
        if isinstance(cell.value, str):
            match = DAY_PATTERN.match(cell.value.strip())
            if match:
                days.append({
                    "day_number": int(match.group(1)),
                    "row": cell.row
                })
    return days

def build_week_day_structure(sheet) -> list[dict]:
    "Combines week detection and day detection into a structured list."
    weeks = find_week_columns(sheet)
    for week in weeks:
        week["days"] = find_day_rows(sheet, week)
    return weeks

def build_all_blocks(workbook) -> list[dict]:
    """Builds the full week/day structure for every training block tab, attaching 
    the sequential block number based on tab order."""
    block_tabs = get_block_tabs(workbook)
    blocks = []
    for block_number, tab_name in enumerate(block_tabs, start=1):
        sheet = workbook[tab_name]
        blocks.append({
            "block_number": block_number,
            "tab_name": tab_name,
            "weeks": build_week_day_structure(sheet)
        })
    return blocks

def resolve_day_end(sheet, day: dict, search_limit: int, exercise_col: int) -> None:
    """Scans downward from day['start_row'] until a blank row is found,
    but never searches past search_limit (the next day's row, or sheet's max_row)."""
    row = day["start_row"]
    while row < search_limit:
        cell = sheet.cell(row=row, column=exercise_col)
        if cell.value is None:
            day["end_row"] = row
            return
        row += 1
    day["end_row"] = search_limit  # reached the limit without finding a blank row

def compute_day_ranges(sheet, week: dict) -> list[dict]:
    """Adds start_row and end_row to each day, using blank-row detection bounded
    by the next day's row (or sheet end, for the last day)."""
    sheet.calculate_dimension(force=True)  # Ensure the sheet's dimensions are up-to-date
    days = week["days"]
    exercise_col = week["start_column"]
    max_row = sheet.max_row
    for i, day in enumerate(days):
        day["start_row"] = day["row"] + 1  # Exercises start the row after "Day N"
        search_limit = days[i + 1]["row"] if i + 1 < len(days) else max_row + 1
        resolve_day_end(sheet, day, search_limit, exercise_col)
    return days

def parse_day_headers(sheet, week: dict, day: dict) -> dict:
    """Reads the header row for a given day and returns a dict mapping
    field name -> column number, handling the duplicate 'load' header case."""
    header_row = day["row"]
    start_col = week["start_column"]
    max_col = start_col + 10  # generous upper bound for one day's column band; adjust if needed

    load_columns = []
    field_columns = {}

    for col in range(start_col, max_col):
        cell = sheet.cell(row=header_row, column=col)
        if not isinstance(cell.value, str):
            continue
        header = cell.value.strip().lower()
        if header == "load":
            load_columns.append(col)
        elif header in ("rpe", "notes", "range", "top end"):
            field_columns[header] = col

    if len(load_columns) >= 2:
        field_columns["prescribed_rpe_or_backoff"] = load_columns[0]
        field_columns["actual_load"] = load_columns[1]
    elif len(load_columns) == 1:
        field_columns["actual_load"] = load_columns[0]  # only one load column present — no prescribed value to disambiguate

    return field_columns

def extract_day_exercises(sheet, week: dict, day: dict) -> list[dict]:
    """Extracts each exercise row within a day's row range into a raw dict."""
    fields = parse_day_headers(sheet, week, day)
    name_col = week["start_column"]
    sets_reps_col = week["start_column"] + 1

    exercises = []
    for row in range(day["start_row"], day["end_row"]):
        name_cell = sheet.cell(row=row, column=name_col)
        if name_cell.value is None:
            continue  # skip any stray blank row within the range

        exercise = {
            "block_number": None,  # filled in by caller, which knows the block
            "week_number": week["week_number"],
            "day_number": day["day_number"],
            "row": row,
            "exercise_name": name_cell.value,
            "sets_x_reps": sheet.cell(row=row, column=sets_reps_col).value,
            "prescribed_raw": sheet.cell(row=row, column=fields["prescribed_rpe_or_backoff"]).value if "prescribed_rpe_or_backoff" in fields else None,
            "actual_load_raw": sheet.cell(row=row, column=fields["actual_load"]).value if "actual_load" in fields else None,
            "actual_rpe_raw": sheet.cell(row=row, column=fields["rpe"]).value if "rpe" in fields else None,
            "notes": sheet.cell(row=row, column=fields["notes"]).value if "notes" in fields else None,
        }
        exercises.append(exercise)
    return exercises

if __name__ == "__main__":
    wb = load_client_workbook("Jon Program.xlsx")
    block_tabs = get_block_tabs(wb)
    sheet = wb[block_tabs[0]]
    week1 = find_week_columns(sheet)[0]
    week1["days"] = find_day_rows(sheet, week1)
    compute_day_ranges(sheet, week1)
    day1_exercises = extract_day_exercises(sheet, week1, week1["days"][0])
    for ex in day1_exercises:
        print(ex)
        
    # block_tabs = get_block_tabs(wb)
    # week1 = find_week_columns(wb[block_tabs[0]])[0]
    # week1["days"] = find_day_rows(wb[block_tabs[0]], week1)
    # compute_day_ranges(wb[block_tabs[0]], week1)
    # print(week1["days"])

    # all_blocks = build_all_blocks(wb)
    # print(f"Found {len(all_blocks)} blocks")
    # print(all_blocks[0])

    # block_tabs = get_block_tabs(wb)
    # weeks = find_week_columns(wb[block_tabs[0]])
    # days = find_day_rows(wb[block_tabs[0]], weeks[0])
    # weeks[0]["days"] = days
    # compute_day_ranges(weeks[0])
    # print(weeks[0]["days"])

    # INDIVIDUAL FUNCTION TESTS
    # -----------------------------------------------------------------------------
    # block_tabs = get_block_tabs(wb)
    # print(block_tabs)
    # print_sheet_grid(wb, block_tabs[0]) # inspect the first block tab
    # weeks = find_week_columns(wb[block_tabs[0]])
    # days = find_day_rows(wb[block_tabs[0]], weeks[0])
    # structure = build_week_day_structure(wb[block_tabs[0]])
    